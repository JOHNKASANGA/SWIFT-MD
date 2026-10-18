import io
import re
import zipfile
from typing import Optional

import httpx
from pypdf import PdfReader

try:
    from docx import Document as DocxDocument

    DOCX_SUPPORTED = True
except ImportError:
    DOCX_SUPPORTED = False

try:
    from pptx import Presentation

    PPTX_SUPPORTED = True
except ImportError:
    PPTX_SUPPORTED = False

try:
    import fitz

    PYMUPDF_SUPPORTED = True
except ImportError:
    PYMUPDF_SUPPORTED = False


SKIP_EXTENSIONS = [
    ".mp4",
    ".mp3",
    ".wav",
    ".avi",
    ".mov",
    ".zip",
    ".rar",
]


def is_skippable(url: str) -> bool:
    url_lower = url.lower().split("?")[0]
    return any(url_lower.endswith(extension) for extension in SKIP_EXTENSIONS)


def get_direct_url(file_url: str) -> str:
    """Resolve MediaFire share pages to a direct download URL."""
    if "mediafire.com" not in file_url:
        return file_url

    response = httpx.get(file_url, timeout=30.0, follow_redirects=True)
    response.raise_for_status()

    match = re.search(
        r'href="(https://download\d+\.mediafire\.com/[^"]+)"',
        response.text,
    )

    if not match:
        raise ValueError("MediaFire did not provide a direct download URL.")

    return match.group(1)


def download_file(file_url: str) -> tuple[bytes, str]:
    """Download a public material and return its bytes with its resolved URL."""
    if is_skippable(file_url):
        raise ValueError("This file type is not supported for classification.")

    direct_url = get_direct_url(file_url)
    timeout = 60.0 if "drive.google.com" in direct_url or "usercontent" in direct_url else 30.0

    response = httpx.get(
        direct_url,
        timeout=timeout,
        follow_redirects=True,
    )
    response.raise_for_status()

    if (
        b"Google Drive - Virus scan warning" in response.content
        or "confirm=" in str(response.url)
    ):
        confirm_match = re.search(r'confirm=([^&"]+)', str(response.url))
        file_id_match = re.search(r"id=([^&]+)", direct_url)

        if confirm_match and file_id_match:
            direct_url = (
                "https://drive.usercontent.google.com/download"
                f"?id={file_id_match.group(1)}&export=download"
                f"&confirm={confirm_match.group(1)}&authuser=0"
            )
            response = httpx.get(
                direct_url,
                timeout=60.0,
                follow_redirects=True,
            )
            response.raise_for_status()

    return response.content, direct_url


def detect_document_type(
    file_bytes: bytes,
    file_url: str,
    direct_url: str,
) -> str:
    """Identify a document from its actual bytes before trusting its URL."""
    if file_bytes.startswith(b"%PDF"):
        return "pdf"

    if file_bytes.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                names = archive.namelist()

            if any(name.startswith("word/") for name in names):
                return "docx"

            if any(name.startswith("ppt/") for name in names):
                return "pptx"

            if any(name.startswith("xl/") for name in names):
                return "xlsx"
        except zipfile.BadZipFile:
            pass

    if file_bytes.startswith((b"\x89PNG", b"\xff\xd8\xff")):
        return "image"

    url_lower = f"{file_url}?{direct_url}".lower()

    if ".docx" in url_lower:
        return "docx"

    if ".pptx" in url_lower:
        return "pptx"

    if ".pdf" in url_lower:
        return "pdf"

    return "unknown"


def extract_text_from_pdf(
    file_bytes: bytes,
    max_pages: Optional[int] = None,
    max_chars: Optional[int] = None,
) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = reader.pages[:max_pages] if max_pages else reader.pages
        parts = []

        for page in pages:
            page_text = page.extract_text()

            if page_text and page_text.strip():
                parts.append(page_text.strip())

            combined = "\n\n".join(parts)

            if max_chars and len(combined) >= max_chars:
                return combined[:max_chars]

        text = "\n\n".join(parts)
        return text[:max_chars] if max_chars else text
    except Exception:
        return ""


def extract_document_batch_from_url(
    file_url: str,
    start_page: int = 0,
    page_count: int = 12,
    max_chars: int = 70000,
) -> dict:
    """Extract one bounded, resumable section of a public study document.

    A full textbook must never keep a web request open while hundreds of
    pages are decoded.  PDF pages are therefore processed by index.  Empty
    PDF pages are reported separately so the caller can route them to OCR.
    """
    try:
        file_bytes, direct_url = download_file(file_url)
    except Exception as error:
        return {
            "text": "",
            "file_type": "unknown",
            "total_pages": 0,
            "start_page": start_page,
            "next_page": start_page,
            "complete": False,
            "scanned_pages": [],
            "error": safe_error_message(error),
        }

    file_type = detect_document_type(file_bytes, file_url, direct_url)

    if file_type != "pdf":
        extractors = {
            "docx": extract_text_from_docx,
            "pptx": extract_text_from_pptx,
        }
        extractor = extractors.get(file_type)
        text = extractor(file_bytes, max_chars=max_chars) if extractor else ""
        return {
            "text": text,
            "file_type": file_type,
            "total_pages": 1,
            "start_page": 0,
            "next_page": 1,
            "complete": True,
            "scanned_pages": [],
            "error": None if text.strip() else f"Unsupported or unreadable {file_type} document.",
        }

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        total_pages = len(reader.pages)
        start_page = max(0, min(start_page, total_pages))
        end_page = min(total_pages, start_page + max(1, page_count))
        parts = []
        scanned_pages = []

        for page_number in range(start_page, end_page):
            page_text = reader.pages[page_number].extract_text() or ""
            page_text = page_text.strip()
            if page_text:
                parts.append(f"[Page {page_number + 1}]\n{page_text}")
            else:
                scanned_pages.append(page_number)

            if sum(len(part) for part in parts) >= max_chars:
                break

        return {
            "text": "\n\n".join(parts)[:max_chars],
            "file_type": "pdf",
            "total_pages": total_pages,
            "start_page": start_page,
            "next_page": end_page,
            "complete": end_page >= total_pages,
            "scanned_pages": scanned_pages,
            "error": None,
        }
    except Exception as error:
        return {
            "text": "",
            "file_type": "pdf",
            "total_pages": 0,
            "start_page": start_page,
            "next_page": start_page,
            "complete": False,
            "scanned_pages": [],
            "error": safe_error_message(error),
        }


def safe_error_message(error: Exception) -> str:
    return str(error).strip()[:500] or error.__class__.__name__


def extract_text_from_docx(
    file_bytes: bytes,
    max_chars: Optional[int] = None,
) -> str:
    if not DOCX_SUPPORTED:
        return ""

    try:
        document = DocxDocument(io.BytesIO(file_bytes))
        text = "\n\n".join(
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        return text[:max_chars] if max_chars else text
    except Exception:
        return ""


def extract_text_from_pptx(
    file_bytes: bytes,
    max_chars: Optional[int] = None,
) -> str:
    if not PPTX_SUPPORTED:
        return ""

    try:
        presentation = Presentation(io.BytesIO(file_bytes))
        parts = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", "")

                if text and text.strip():
                    parts.append(text.strip())

                combined = "\n\n".join(parts)

                if max_chars and len(combined) >= max_chars:
                    return combined[:max_chars]

        text = "\n\n".join(parts)
        return text[:max_chars] if max_chars else text
    except Exception:
        return ""


def render_first_pdf_page(file_bytes: bytes) -> bytes:
    """Render the first page of an image-only PDF for vision OCR."""
    if not PYMUPDF_SUPPORTED:
        return b""

    document = None

    try:
        document = fitz.open(stream=file_bytes, filetype="pdf")
        page = document.load_page(0)
        image = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        return image.tobytes("png")
    except Exception:
        return b""
    finally:
        if document is not None:
            document.close()


def extract_material_sample_from_url(
    file_url: str,
    max_pages: int = 4,
    max_chars: int = 14000,
) -> dict:
    """
    Extract enough real material content for classification.

    Text-based documents return text. Image-only PDFs return a rendered first
    page for vision OCR. Unsupported or inaccessible files include an error.
    """
    try:
        file_bytes, direct_url = download_file(file_url)
    except Exception as error:
        return {
            "text": "",
            "image_bytes": b"",
            "file_type": "unknown",
            "error": str(error),
        }

    file_type = detect_document_type(file_bytes, file_url, direct_url)

    if file_type == "pdf":
        text = extract_text_from_pdf(
            file_bytes,
            max_pages=max_pages,
            max_chars=max_chars,
        )

        if text.strip():
            return {
                "text": text,
                "image_bytes": b"",
                "file_type": "pdf",
                "error": None,
            }

        image_bytes = render_first_pdf_page(file_bytes)

        return {
            "text": "",
            "image_bytes": image_bytes,
            "file_type": "scanned_pdf",
            "error": (
                None
                if image_bytes
                else "PDF has no extractable text and could not be rendered for OCR."
            ),
        }

    if file_type == "docx":
        text = extract_text_from_docx(file_bytes, max_chars=max_chars)
    elif file_type == "pptx":
        text = extract_text_from_pptx(file_bytes, max_chars=max_chars)
    else:
        text = ""

    return {
        "text": text,
        "image_bytes": b"",
        "file_type": file_type,
        "error": None if text.strip() else f"Unsupported or unreadable {file_type} document.",
    }


def extract_text_from_url(file_url: str) -> str:
    """Compatibility helper used by Swift's existing text-extraction routes."""
    sample = extract_material_sample_from_url(
        file_url,
        max_pages=1000,
        max_chars=None,
    )
    return sample["text"]


def chunk_text(
    text: str,
    chunk_size: int = 4000,
    overlap: int = 200,
) -> list[str]:
    """Split text into overlapping chunks for better question coverage."""
    if not text or len(text) <= chunk_size:
        return [text] if text else []

    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size

        if end < len(text):
            boundary = text.rfind(".", start + chunk_size - 500, end)

            if boundary > start:
                end = boundary + 1

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start = end - overlap

        if start >= len(text):
            break

    return chunks
